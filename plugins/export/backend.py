"""TXT / DOCX 导出插件:数据访问 + 路由注册(原 services/export_service.py 与 app.py Export 路由迁移而来)

无专属表(直接 SQL 读 chapters/drafts/volumes/acts/graph_nodes 等表),故无 migrations 目录。
对外 provide "export" 服务:json_transfer 插件的整项目 JSON 导出复用 resolve_output_path,
运行时通过 api.require("export") 获取,避免插件间直接 import。
"""
import json
from pathlib import Path

_api = None  # activate 时注入的 PluginAPI

def _resolve_output_path(output_path, ext, default_name):
    """解析导出路径：留空→桌面；目录→放入默认文件名；缺扩展名→自动补全"""
    p = (output_path or "").strip()
    if not p:
        p = str(Path.home() / "Desktop" / f"{default_name}.{ext}")
    path = Path(p)
    if path.is_dir():
        path = path / f"{default_name}.{ext}"
    if path.suffix.lower() != f".{ext}":
        path = path.with_suffix(f".{ext}")
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        raise ValueError(f"无法创建目录 {path.parent}：{e}")
    return str(path)

def _get_volume_preface(conn, vol):
    """取卷首语：volumes.preface 为空时，回退到 drafts 表该卷的当前卷首语草稿"""
    if vol["preface"]:
        return vol["preface"]
    row = conn.execute(
        "SELECT content FROM drafts WHERE volume_id = ? AND chapter_id IS NULL AND is_current = 1 ORDER BY version_number DESC LIMIT 1",
        (vol["id"],)
    ).fetchone()
    return row["content"] if row and row["content"] else ""

def _get_chapter_text(conn, chapter_id):
    """取章节正文：正文已节点化(chapter 插件注册的 text 节点,payload.content),
    一律过 payload 条件按 chapter_id 取(异构父约定,id 序列会撞,不能只信 parent_id 列);
    节点缺失时回退 drafts 当前版本(未迁移/异常状态的兜底,保证导出内容不丢)"""
    node = conn.execute(
        """SELECT json_extract(payload, '$.content') AS content FROM graph_nodes
           WHERE type = 'text' AND json_extract(payload, '$.chapter_id') = ?
             AND deleted_at IS NULL ORDER BY id LIMIT 1""",
        (chapter_id,)
    ).fetchone()
    if node:
        return node["content"] or ""
    draft = conn.execute(
        "SELECT content FROM drafts WHERE chapter_id = ? AND is_current = 1 ORDER BY version_number DESC LIMIT 1",
        (chapter_id,)
    ).fetchone()
    return draft["content"] if draft and draft["content"] else ""


def _get_chapter_content(chapter_id):
    conn = _api.db()
    ch = conn.execute("SELECT * FROM chapters WHERE id = ? AND deleted_at IS NULL", (chapter_id,)).fetchone()
    if not ch: conn.close(); return None
    vol = conn.execute("SELECT * FROM volumes WHERE id = ? AND deleted_at IS NULL", (ch["volume_id"],)).fetchone()
    content = _get_chapter_text(conn, chapter_id)
    # 章节在所属卷内的序号（未删除章节中按 sort_order 排第几）
    idx_row = conn.execute(
        "SELECT COUNT(*) + 1 AS n FROM chapters WHERE volume_id = ? AND deleted_at IS NULL AND sort_order < ?",
        (ch["volume_id"], ch["sort_order"])
    ).fetchone()
    preface = _get_volume_preface(conn, vol) if vol else ""
    conn.close()
    return {
        "volumeTitle": vol["title"] if vol else "",
        "volumePreface": preface,
        "chapterTitle": ch["title"],
        "chapterIndex": idx_row["n"],
        "content": content,
    }

def _get_project_contents(project_id):
    conn = _api.db()
    volumes = conn.execute("SELECT * FROM volumes WHERE project_id = ? AND deleted_at IS NULL ORDER BY sort_order", (project_id,)).fetchall()
    # 正文节点一次取回(text 节点,payload.content;过 payload 条件,不信 parent_id 列)
    text_map = {r["chapter_id"]: (r["content"] or "") for r in conn.execute(
        """SELECT json_extract(payload, '$.chapter_id') AS chapter_id,
                  json_extract(payload, '$.content') AS content
           FROM graph_nodes
           WHERE project_id = ? AND type = 'text' AND deleted_at IS NULL ORDER BY id""",
        (project_id,)).fetchall() if r["chapter_id"] is not None}
    result = []
    for vol in volumes:
        chapters = conn.execute("SELECT * FROM chapters WHERE volume_id = ? AND deleted_at IS NULL ORDER BY sort_order", (vol["id"],)).fetchall()
        ch_contents = []
        for ch_idx, ch in enumerate(chapters):
            # 节点缺失时回退 drafts 当前版本(兜底,见 _get_chapter_text)
            content = text_map[ch["id"]] if ch["id"] in text_map else _get_chapter_text(conn, ch["id"])
            ch_contents.append({"index": ch_idx + 1, "title": ch["title"], "content": content})
        result.append({
            "volumeTitle": vol["title"],
            "volumePreface": _get_volume_preface(conn, vol),
            "chapters": ch_contents,
        })
    conn.close()
    return result


def _get_project_type(project_id):
    conn = _api.db()
    row = conn.execute("SELECT project_type, title FROM projects WHERE id = ?", (project_id,)).fetchone()
    conn.close()
    return (row["project_type"], row["title"]) if row else ("novel", "export")


def _node_to_element(row):
    """graph_nodes 行展开为导出用的元素 dict(剧作元素已迁入 graph_nodes,见 script 插件)。
    只取导出排版需要的字段:节点 type 即 element_type,内容字段在 payload"""
    try:
        p = json.loads(row["payload"] or "{}")
    except ValueError:
        p = {}
    return {
        "id": row["id"],
        "element_type": row["type"],
        "character_id": p.get("character_id"),
        "content": p.get("content"),
        "song_title": p.get("song_title"),
        "score_file": p.get("score_file"),
    }


def _get_script_contents(project_id):
    """剧本/音乐剧结构：幕 → 场 → 元素（容器可嵌套：歌曲 → 重唱 → 分部唱词）。
    元素存 graph_nodes:按场/按父元素取数一律过 payload 条件(scene_id/parent_id),
    不能只看 parent_id 列(异构父约定:顶层元素的 parent_id 列是场景 id,与节点 id 可能数值相撞)"""
    conn = _api.db()
    char_names = {r["id"]: r["name"] for r in conn.execute(
        "SELECT id, name FROM characters WHERE project_id = ?", (project_id,)).fetchall()}

    def build_element(row, depth=0):
        e = _node_to_element(row)
        id_rows = conn.execute(
            "SELECT character_id FROM element_characters WHERE element_id = ? ORDER BY sort_order, id",
            (e["id"],)).fetchall()
        ids = [r["character_id"] for r in id_rows]
        if not ids and e["character_id"]:
            ids = [e["character_id"]]
        names = [char_names.get(cid, "") for cid in ids]
        e["character_ids"] = ids
        name_str = "、".join(n for n in names if n)
        # 齐白：多个角色同说一句对白，标注（齐）
        if e["element_type"] == "dialogue" and len(ids) > 1 and name_str:
            name_str += "（齐）"
        e["characterName"] = name_str
        if e["element_type"] in ("song", "ensemble", "dual") and depth < 4:
            children = conn.execute(
                """SELECT * FROM graph_nodes WHERE json_extract(payload, '$.parent_id') = ?
                   AND deleted_at IS NULL ORDER BY sort_order""", (e["id"],)).fetchall()
            e["children"] = [build_element(c, depth + 1) for c in children]
        return e

    acts = conn.execute(
        "SELECT * FROM acts WHERE project_id = ? AND deleted_at IS NULL ORDER BY sort_order", (project_id,)).fetchall()
    result = []
    for act in acts:
        scenes = conn.execute(
            "SELECT * FROM scenes WHERE act_id = ? AND deleted_at IS NULL ORDER BY sort_order", (act["id"],)).fetchall()
        sc_list = []
        for sc in scenes:
            elems = conn.execute(
                """SELECT * FROM graph_nodes WHERE json_extract(payload, '$.scene_id') = ?
                   AND json_extract(payload, '$.parent_id') IS NULL
                   AND deleted_at IS NULL ORDER BY sort_order""", (sc["id"],)).fetchall()
            sc_list.append({
                "title": sc["title"],
                "setting": sc["setting"] or "",
                "elements": [build_element(el) for el in elems],
            })
        result.append({"actTitle": act["title"], "scenes": sc_list})
    conn.close()
    return result


def _merge_action_dialogue(elements):
    """同一角色的「动作 + 对白」相邻时合并为一项，导出时排成一段"""
    merged = []
    i = 0
    while i < len(elements):
        e = elements[i]
        if e["element_type"] == "action" and e.get("characterName") and i + 1 < len(elements):
            nxt = elements[i + 1]
            if nxt["element_type"] == "dialogue" and nxt.get("characterName") == e["characterName"]:
                merged.append({**e, "element_type": "action_dialogue",
                               "dialogue_content": nxt["content"]})
                i += 2
                continue
        merged.append(e)
        i += 1
    return merged


def _format_element_txt(e, indent=""):
    t = e["element_type"]
    who = f"{e['characterName']}" if e.get("characterName") else ""
    if t == "action_dialogue":
        return f"{indent}{who}（动作）{e['content']} / {e['dialogue_content']}"
    if t == "action":
        return f"{indent}{who}（动作）{e['content']}"
    if t == "dialogue":
        return f"{indent}{who}{e['content']}"
    if t == "lyric":
        lines = (e["content"] or "").split("\n")
        return f"{indent}{who}（唱）\n" + "\n".join(f"{indent}    {l}" for l in lines)
    if t == "ensemble":
        lines = [f"{indent}（重唱）"]
        for c in e.get("children", []):
            part = _format_element_txt(c, indent + "    ")
            lines.append(part)
        return "\n".join(lines)
    if t == "dual":
        # 叠白：多人同时说不同的话，逐栏缩进列出
        lines = [f"{indent}（叠白）"]
        for c in e.get("children", []):
            lines.append(_format_element_txt(c, indent + "    "))
        return "\n".join(lines)
    if t == "song":
        score_mark = "（附乐谱）" if e.get("score_file") else ""
        lines = [f"{indent}♪ {e['song_title'] or '（未命名歌曲）'}{score_mark}"]
        for c in e.get("children", []):
            lines.append(_format_element_txt(c, indent + "    "))
        lines.append(f"{indent}（歌曲结束）")
        return "\n".join(lines)
    return f"{indent}{e['content']}"

def export_txt(chapter_id=None, project_id=None, output_path=""):
    text = ""
    default_name = "export"
    if chapter_id:
        data = _get_chapter_content(chapter_id)
        if not data: raise ValueError("Chapter not found")
        default_name = data["chapterTitle"] or default_name
        if data["volumeTitle"]: text += f"【{data['volumeTitle']}】\n\n"
        text += f"第{data['chapterIndex']}章 {data['chapterTitle']}\n\n{data['content']}"
    elif project_id:
        ptype, ptitle = _get_project_type(project_id)
        default_name = ptitle or default_name
        if ptype == "novel":
            volumes = _get_project_contents(project_id)
            for vol in volumes:
                text += f"\n\n========== {vol['volumeTitle']} ==========\n\n"
                if vol["volumePreface"]: text += f"{vol['volumePreface']}\n\n---\n\n"
                for ch in vol["chapters"]:
                    text += f"第{ch['index']}章 {ch['title']}\n\n{ch['content']}\n\n"
        else:
            for act in _get_script_contents(project_id):
                text += f"\n\n========== {act['actTitle']} ==========\n\n"
                for sc in act["scenes"]:
                    text += f"【{sc['title']}】\n"
                    if sc["setting"]: text += f"（场景）{sc['setting']}\n"
                    text += "\n"
                    for e in _merge_action_dialogue(sc["elements"]):
                        text += _format_element_txt(e) + "\n\n"
    out = _resolve_output_path(output_path, "txt", default_name)
    try:
        with open(out, "w", encoding="utf-8") as f:
            f.write(text)
    except OSError as e:
        raise ValueError(f"无法写入文件 {out}：{e}")
    return {"filePath": out}


def export_docx(chapter_id=None, project_id=None, output_path="", options=None):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    opts = options or {}
    font_name = opts.get("font", "SimSun")
    font_size = opts.get("fontSize", 12)  # 12pt
    line_spacing = opts.get("lineSpacing", 1.5)

    # 东亚字体名映射（避免 Word 回退到 MS Gothic / ＭＳ 明朝）
    EA_FONTS = {
        "SimSun": "宋体", "SimHei": "黑体", "KaiTi": "楷体",
        "FangSong": "仿宋", "Microsoft YaHei": "微软雅黑",
    }
    ea_body = EA_FONTS.get(font_name, "宋体")
    ASCII_FONT = "Times New Roman"  # 所有非角色名西文统一 TNR

    def set_ea_style(style, ea_name):
        style.font.name = ASCII_FONT
        rPr = style.element.get_or_add_rPr()
        rPr.get_or_add_rFonts().set(qn("w:eastAsia"), ea_name)

    def style_run(run, ea_name, bold=False, italic=False, ascii_font=None):
        run.font.name = ascii_font or ASCII_FONT
        rPr = run._element.get_or_add_rPr()
        rPr.get_or_add_rFonts().set(qn("w:eastAsia"), ea_name)
        if bold: run.font.bold = True
        if italic: run.font.italic = True
        return run

    def add_text_with_breaks(p, text, ea_name, bold=False, italic=False):
        """多行文本写入同一段落，行间用软换行"""
        lines = (text or "").split("\n")
        for i, line in enumerate(lines):
            run = style_run(p.add_run(line), ea_name, bold=bold, italic=italic)
            if i < len(lines) - 1:
                run.add_break()

    def hanging_indent(p, chars=4):
        """悬挂缩进：左缩进 chars 字符 + 悬挂 chars 字符"""
        pPr = p._element.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        ind.set(qn("w:leftChars"), str(chars * 100))
        ind.set(qn("w:hangingChars"), str(chars * 100))

    def left_indent(p, chars):
        pPr = p._element.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        ind.set(qn("w:leftChars"), str(chars * 100))

    def add_script_element(e, in_song=False):
        """排版：角色名黑体加粗（无冒号）、动作楷体、唱词仿宋，悬挂缩进4字符；
        歌曲内：（唱）/（念白）标记后换行；歌曲结束加（歌曲结束）标识"""
        t = e["element_type"]
        name = e.get("characterName") or ""
        if t == "action_dialogue":
            # 同一角色的动作+对白合并段：角色名黑体加粗，动作楷体（西文斜体 TNR），对白正文宋体
            p = doc.add_paragraph()
            hanging_indent(p, 4)
            style_run(p.add_run(name), "黑体", bold=True, ascii_font="黑体")
            add_text_with_breaks(p, e["content"], "楷体", italic=True)
            style_run(p.add_run(" "), ea_body)
            add_text_with_breaks(p, e["dialogue_content"], ea_body)
        elif t == "action":
            p = doc.add_paragraph()
            if name:
                style_run(p.add_run(name), "黑体", bold=True, ascii_font="黑体")
            style_run(p.add_run("（动作）"), "楷体")
            add_text_with_breaks(p, e["content"], "楷体")
        elif t == "dialogue":
            p = doc.add_paragraph()
            hanging_indent(p, 4)
            if name:
                style_run(p.add_run(name), "黑体", bold=True, ascii_font="黑体")
            if in_song:
                # 歌曲内对白视为念白，标记后换行
                style_run(p.add_run("（念白）"), "黑体", bold=True, ascii_font="黑体").add_break()
            add_text_with_breaks(p, e["content"], ea_body)
        elif t == "lyric":
            p = doc.add_paragraph()
            hanging_indent(p, 4)
            if name:
                style_run(p.add_run(name), "黑体", bold=True, ascii_font="黑体")
            # （唱）标记后换行
            style_run(p.add_run("（唱）"), "黑体", bold=True, ascii_font="黑体").add_break()
            add_text_with_breaks(p, e["content"], "仿宋")
        elif t == "ensemble":
            # 重唱：无边框表格，每个角色一列并排显示
            p = doc.add_paragraph()
            style_run(p.add_run("（重唱）"), "黑体", bold=True)
            children = e.get("children", [])
            if children:
                table = doc.add_table(rows=1, cols=len(children))
                for i, c in enumerate(children):
                    cp = table.rows[0].cells[i].paragraphs[0]
                    cname = c.get("characterName") or ""
                    if cname:
                        style_run(cp.add_run(cname), "黑体", bold=True, ascii_font="黑体")
                    style_run(cp.add_run("（唱）"), "黑体", bold=True, ascii_font="黑体").add_break()
                    lines = (c["content"] or "").split("\n")
                    for j, line in enumerate(lines):
                        run = style_run(cp.add_run(line), "仿宋")
                        if j < len(lines) - 1:
                            run.add_break()
        elif t == "dual":
            # 叠白：无边框表格，每个角色一列并排显示（同时对白）
            p = doc.add_paragraph()
            style_run(p.add_run("（叠白）"), "黑体", bold=True)
            children = e.get("children", [])
            if children:
                table = doc.add_table(rows=1, cols=len(children))
                for i, c in enumerate(children):
                    cp = table.rows[0].cells[i].paragraphs[0]
                    cname = c.get("characterName") or ""
                    if cname:
                        style_run(cp.add_run(cname), "黑体", bold=True, ascii_font="黑体").add_break()
                    lines = (c["content"] or "").split("\n")
                    for j, line in enumerate(lines):
                        run = style_run(cp.add_run(line), ea_body)
                        if j < len(lines) - 1:
                            run.add_break()
        elif t == "song":
            p = doc.add_paragraph()
            score_mark = "（附乐谱）" if e.get("score_file") else ""
            style_run(p.add_run(f"♪ {e['song_title'] or '（未命名歌曲）'}{score_mark}"), "黑体", bold=True)
            for c in e.get("children", []):
                add_script_element(c, in_song=True)
            # 歌曲结束标识
            end_p = doc.add_paragraph()
            style_run(end_p.add_run("（歌曲结束）"), "黑体", bold=True)
        else:
            p = doc.add_paragraph()
            add_text_with_breaks(p, e["content"], ea_body)

    doc = Document()
    style = doc.styles["Normal"]
    set_ea_style(style, ea_body)
    style.font.size = Pt(font_size)
    style.paragraph_format.line_spacing = line_spacing
    # 标题样式同样显式指定中文字体，避免主题字体回退到日文字体
    for hs in ["Heading 1", "Heading 2", "Heading 3", "Title"]:
        try:
            set_ea_style(doc.styles[hs], "黑体")
        except KeyError:
            pass

    default_name = "export"
    if chapter_id:
        data = _get_chapter_content(chapter_id)
        if not data: raise ValueError("Chapter not found")
        default_name = data["chapterTitle"] or default_name
        if data["volumeTitle"]:
            doc.add_heading(data["volumeTitle"], level=1)
        doc.add_heading(data["chapterTitle"], level=2)
        for para in (data["content"] or "").split("\n"):
            doc.add_paragraph(para)
    elif project_id:
        ptype, ptitle = _get_project_type(project_id)
        default_name = ptitle or default_name
        if ptype == "novel":
            volumes = _get_project_contents(project_id)
            for vol in volumes:
                doc.add_heading(vol["volumeTitle"], level=1)
                if vol["volumePreface"]:
                    # 斜体需设置在 run 上（python-docx 的 Paragraph 没有 italic 属性）
                    p = doc.add_paragraph()
                    style_run(p.add_run(vol["volumePreface"]), ea_body, italic=True)
                for ch in vol["chapters"]:
                    doc.add_heading(ch["title"], level=2)
                    for para in (ch["content"] or "").split("\n"):
                        doc.add_paragraph(para)
        else:
            for act in _get_script_contents(project_id):
                doc.add_heading(act["actTitle"], level=1)
                for sc in act["scenes"]:
                    doc.add_heading(sc["title"], level=2)
                    if sc["setting"]:
                        p = doc.add_paragraph()
                        style_run(p.add_run(f"（场景）{sc['setting']}"), "楷体", italic=True)
                    for e in _merge_action_dialogue(sc["elements"]):
                        add_script_element(e)

    out = _resolve_output_path(output_path, "docx", default_name)
    try:
        doc.save(out)
    except OSError as e:
        raise ValueError(f"无法写入文件 {out}：{e}")
    return {"filePath": out}


class Plugin:
    def activate(self, api):
        global _api
        _api = api

        # 对外提供导出路径解析,json_transfer 插件的整项目 JSON 导出依赖它(运行时再取,避免加载顺序耦合)
        api.provide("export", {
            "resolve_output_path": _resolve_output_path,
        })

        @api.route("/api/export/txt", methods=["POST"])
        def api_export_txt():
            data = api.req_json()
            try:
                return api.jsonify(export_txt(
                    chapter_id=data.get("chapterId"),
                    project_id=data.get("projectId"),
                    output_path=data.get("outputPath", ""),
                ))
            except ValueError as e:
                return api.jsonify({"error": str(e)}), 400

        @api.route("/api/export/docx", methods=["POST"])
        def api_export_docx():
            data = api.req_json()
            try:
                return api.jsonify(export_docx(
                    chapter_id=data.get("chapterId"),
                    project_id=data.get("projectId"),
                    output_path=data.get("outputPath", ""),
                    options=data.get("options"),
                ))
            except ValueError as e:
                return api.jsonify({"error": str(e)}), 400
