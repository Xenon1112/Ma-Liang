import os
from pathlib import Path
from database import get_conn, row_to_dict

def _resolve_output_path(output_path, ext, default_name):
    """解析导出路径：留空→桌面；目录→放入默认文件名；缺扩展名→自动补全"""
    p = (output_path or "").strip()
    if not p:
        p = str(Path.home() / "Desktop" / f"{default_name}.{ext}")
    path = Path(p)
    if path.is_dir():
        path = path / f"{default_name}.{ext}"
    if path.suffix.lower() != f".{ext}":
        path = path.with_suffix(f".{ext}")
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        raise ValueError(f"无法创建目录 {path.parent}：{e}")
    return str(path)

def _get_volume_preface(conn, vol):
    """取卷首语：volumes.preface 为空时，回退到 drafts 表该卷的当前卷首语草稿"""
    if vol["preface"]:
        return vol["preface"]
    row = conn.execute(
        "SELECT content FROM drafts WHERE volume_id = ? AND chapter_id IS NULL AND is_current = 1 ORDER BY version_number DESC LIMIT 1",
        (vol["id"],)
    ).fetchone()
    return row["content"] if row and row["content"] else ""

def _get_chapter_content(chapter_id):
    conn = get_conn()
    ch = conn.execute("SELECT * FROM chapters WHERE id = ? AND deleted_at IS NULL", (chapter_id,)).fetchone()
    if not ch: conn.close(); return None
    vol = conn.execute("SELECT * FROM volumes WHERE id = ? AND deleted_at IS NULL", (ch["volume_id"],)).fetchone()
    draft = conn.execute(
        "SELECT * FROM drafts WHERE chapter_id = ? AND is_current = 1 ORDER BY version_number DESC LIMIT 1",
        (chapter_id,)
    ).fetchone()
    # 章节在所属卷内的序号（未删除章节中按 sort_order 排第几）
    idx_row = conn.execute(
        "SELECT COUNT(*) + 1 AS n FROM chapters WHERE volume_id = ? AND deleted_at IS NULL AND sort_order < ?",
        (ch["volume_id"], ch["sort_order"])
    ).fetchone()
    preface = _get_volume_preface(conn, vol) if vol else ""
    conn.close()
    return {
        "volumeTitle": vol["title"] if vol else "",
        "volumePreface": preface,
        "chapterTitle": ch["title"],
        "chapterIndex": idx_row["n"],
        "content": draft["content"] if draft else "",
    }

def _get_project_contents(project_id):
    conn = get_conn()
    volumes = conn.execute("SELECT * FROM volumes WHERE project_id = ? AND deleted_at IS NULL ORDER BY sort_order", (project_id,)).fetchall()
    result = []
    for vol in volumes:
        chapters = conn.execute("SELECT * FROM chapters WHERE volume_id = ? AND deleted_at IS NULL ORDER BY sort_order", (vol["id"],)).fetchall()
        ch_contents = []
        for ch_idx, ch in enumerate(chapters):
            draft = conn.execute("SELECT * FROM drafts WHERE chapter_id = ? AND is_current = 1 ORDER BY version_number DESC LIMIT 1", (ch["id"],)).fetchone()
            ch_contents.append({"index": ch_idx + 1, "title": ch["title"], "content": draft["content"] if draft else ""})
        result.append({
            "volumeTitle": vol["title"],
            "volumePreface": _get_volume_preface(conn, vol),
            "chapters": ch_contents,
        })
    conn.close()
    return result


def _get_project_type(project_id):
    conn = get_conn()
    row = conn.execute("SELECT project_type, title FROM projects WHERE id = ?", (project_id,)).fetchone()
    conn.close()
    return (row["project_type"], row["title"]) if row else ("novel", "export")


def _get_script_contents(project_id):
    """剧本/音乐剧结构：幕 → 场 → 元素（容器可嵌套：歌曲 → 重唱 → 分部唱词）"""
    conn = get_conn()
    char_names = {r["id"]: r["name"] for r in conn.execute(
        "SELECT id, name FROM characters WHERE project_id = ?", (project_id,)).fetchall()}

    def build_element(row, depth=0):
        e = row_to_dict(row)
        id_rows = conn.execute(
            "SELECT character_id FROM element_characters WHERE element_id = ? ORDER BY sort_order, id",
            (e["id"],)).fetchall()
        ids = [r["character_id"] for r in id_rows]
        if not ids and e["character_id"]:
            ids = [e["character_id"]]
        names = [char_names.get(cid, "") for cid in ids]
        e["character_ids"] = ids
        name_str = "、".join(n for n in names if n)
        # 齐白：多个角色同说一句对白，标注（齐）
        if e["element_type"] == "dialogue" and len(ids) > 1 and name_str:
            name_str += "（齐）"
        e["characterName"] = name_str
        if e["element_type"] in ("song", "ensemble", "dual") and depth < 4:
            children = conn.execute(
                """SELECT * FROM script_elements WHERE parent_id = ? AND deleted_at IS NULL
                   ORDER BY sort_order""", (e["id"],)).fetchall()
            e["children"] = [build_element(c, depth + 1) for c in children]
        return e

    acts = conn.execute(
        "SELECT * FROM acts WHERE project_id = ? AND deleted_at IS NULL ORDER BY sort_order", (project_id,)).fetchall()
    result = []
    for act in acts:
        scenes = conn.execute(
            "SELECT * FROM scenes WHERE act_id = ? AND deleted_at IS NULL ORDER BY sort_order", (act["id"],)).fetchall()
        sc_list = []
        for sc in scenes:
            elems = conn.execute(
                """SELECT * FROM script_elements WHERE scene_id = ? AND parent_id IS NULL
                   AND deleted_at IS NULL ORDER BY sort_order""", (sc["id"],)).fetchall()
            sc_list.append({
                "title": sc["title"],
                "setting": sc["setting"] or "",
                "elements": [build_element(el) for el in elems],
            })
        result.append({"actTitle": act["title"], "scenes": sc_list})
    conn.close()
    return result


def _merge_action_dialogue(elements):
    """同一角色的「动作 + 对白」相邻时合并为一项，导出时排成一段"""
    merged = []
    i = 0
    while i < len(elements):
        e = elements[i]
        if e["element_type"] == "action" and e.get("characterName") and i + 1 < len(elements):
            nxt = elements[i + 1]
            if nxt["element_type"] == "dialogue" and nxt.get("characterName") == e["characterName"]:
                merged.append({**e, "element_type": "action_dialogue",
                               "dialogue_content": nxt["content"]})
                i += 2
                continue
        merged.append(e)
        i += 1
    return merged


def _format_element_txt(e, indent=""):
    t = e["element_type"]
    who = f"{e['characterName']}" if e.get("characterName") else ""
    if t == "action_dialogue":
        return f"{indent}{who}（动作）{e['content']} / {e['dialogue_content']}"
    if t == "action":
        return f"{indent}{who}（动作）{e['content']}"
    if t == "dialogue":
        return f"{indent}{who}{e['content']}"
    if t == "lyric":
        lines = (e["content"] or "").split("\n")
        return f"{indent}{who}（唱）\n" + "\n".join(f"{indent}    {l}" for l in lines)
    if t == "ensemble":
        lines = [f"{indent}（重唱）"]
        for c in e.get("children", []):
            part = _format_element_txt(c, indent + "    ")
            lines.append(part)
        return "\n".join(lines)
    if t == "dual":
        # 叠白：多人同时说不同的话，逐栏缩进列出
        lines = [f"{indent}（叠白）"]
        for c in e.get("children", []):
            lines.append(_format_element_txt(c, indent + "    "))
        return "\n".join(lines)
    if t == "song":
        lines = [f"{indent}♪ {e['song_title'] or '（未命名歌曲）'}"]
        for c in e.get("children", []):
            lines.append(_format_element_txt(c, indent + "    "))
        lines.append(f"{indent}（歌曲结束）")
        return "\n".join(lines)
    return f"{indent}{e['content']}"

def export_txt(chapter_id=None, project_id=None, output_path=""):
    text = ""
    default_name = "export"
    if chapter_id:
        data = _get_chapter_content(chapter_id)
        if not data: raise ValueError("Chapter not found")
        default_name = data["chapterTitle"] or default_name
        if data["volumeTitle"]: text += f"【{data['volumeTitle']}】\n\n"
        text += f"第{data['chapterIndex']}章 {data['chapterTitle']}\n\n{data['content']}"
    elif project_id:
        ptype, ptitle = _get_project_type(project_id)
        default_name = ptitle or default_name
        if ptype == "novel":
            volumes = _get_project_contents(project_id)
            for vol in volumes:
                text += f"\n\n========== {vol['volumeTitle']} ==========\n\n"
                if vol["volumePreface"]: text += f"{vol['volumePreface']}\n\n---\n\n"
                for ch in vol["chapters"]:
                    text += f"第{ch['index']}章 {ch['title']}\n\n{ch['content']}\n\n"
        else:
            for act in _get_script_contents(project_id):
                text += f"\n\n========== {act['actTitle']} ==========\n\n"
                for sc in act["scenes"]:
                    text += f"【{sc['title']}】\n"
                    if sc["setting"]: text += f"（场景）{sc['setting']}\n"
                    text += "\n"
                    for e in _merge_action_dialogue(sc["elements"]):
                        text += _format_element_txt(e) + "\n\n"
    out = _resolve_output_path(output_path, "txt", default_name)
    try:
        with open(out, "w", encoding="utf-8") as f:
            f.write(text)
    except OSError as e:
        raise ValueError(f"无法写入文件 {out}：{e}")
    return {"filePath": out}


def export_docx(chapter_id=None, project_id=None, output_path="", options=None):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    opts = options or {}
    font_name = opts.get("font", "SimSun")
    font_size = opts.get("fontSize", 12)  # 12pt
    line_spacing = opts.get("lineSpacing", 1.5)

    # 东亚字体名映射（避免 Word 回退到 MS Gothic / ＭＳ 明朝）
    EA_FONTS = {
        "SimSun": "宋体", "SimHei": "黑体", "KaiTi": "楷体",
        "FangSong": "仿宋", "Microsoft YaHei": "微软雅黑",
    }
    ea_body = EA_FONTS.get(font_name, "宋体")
    ASCII_FONT = "Times New Roman"  # 所有非角色名西文统一 TNR

    def set_ea_style(style, ea_name):
        style.font.name = ASCII_FONT
        rPr = style.element.get_or_add_rPr()
        rPr.get_or_add_rFonts().set(qn("w:eastAsia"), ea_name)

    def style_run(run, ea_name, bold=False, italic=False, ascii_font=None):
        run.font.name = ascii_font or ASCII_FONT
        rPr = run._element.get_or_add_rPr()
        rPr.get_or_add_rFonts().set(qn("w:eastAsia"), ea_name)
        if bold: run.font.bold = True
        if italic: run.font.italic = True
        return run

    def add_text_with_breaks(p, text, ea_name, bold=False, italic=False):
        """多行文本写入同一段落，行间用软换行"""
        lines = (text or "").split("\n")
        for i, line in enumerate(lines):
            run = style_run(p.add_run(line), ea_name, bold=bold, italic=italic)
            if i < len(lines) - 1:
                run.add_break()

    def hanging_indent(p, chars=4):
        """悬挂缩进：左缩进 chars 字符 + 悬挂 chars 字符"""
        pPr = p._element.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        ind.set(qn("w:leftChars"), str(chars * 100))
        ind.set(qn("w:hangingChars"), str(chars * 100))

    def left_indent(p, chars):
        pPr = p._element.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        ind.set(qn("w:leftChars"), str(chars * 100))

    def add_script_element(e, in_song=False):
        """排版：角色名黑体加粗（无冒号）、动作楷体、唱词仿宋，悬挂缩进4字符；
        歌曲内：（唱）/（念白）标记后换行；歌曲结束加（歌曲结束）标识"""
        t = e["element_type"]
        name = e.get("characterName") or ""
        if t == "action_dialogue":
            # 同一角色的动作+对白合并段：角色名黑体加粗，动作楷体（西文斜体 TNR），对白正文宋体
            p = doc.add_paragraph()
            hanging_indent(p, 4)
            style_run(p.add_run(name), "黑体", bold=True, ascii_font="黑体")
            add_text_with_breaks(p, e["content"], "楷体", italic=True)
            style_run(p.add_run(" "), ea_body)
            add_text_with_breaks(p, e["dialogue_content"], ea_body)
        elif t == "action":
            p = doc.add_paragraph()
            if name:
                style_run(p.add_run(name), "黑体", bold=True, ascii_font="黑体")
            style_run(p.add_run("（动作）"), "楷体")
            add_text_with_breaks(p, e["content"], "楷体")
        elif t == "dialogue":
            p = doc.add_paragraph()
            hanging_indent(p, 4)
            if name:
                style_run(p.add_run(name), "黑体", bold=True, ascii_font="黑体")
            if in_song:
                # 歌曲内对白视为念白，标记后换行
                style_run(p.add_run("（念白）"), "黑体", bold=True, ascii_font="黑体").add_break()
            add_text_with_breaks(p, e["content"], ea_body)
        elif t == "lyric":
            p = doc.add_paragraph()
            hanging_indent(p, 4)
            if name:
                style_run(p.add_run(name), "黑体", bold=True, ascii_font="黑体")
            # （唱）标记后换行
            style_run(p.add_run("（唱）"), "黑体", bold=True, ascii_font="黑体").add_break()
            add_text_with_breaks(p, e["content"], "仿宋")
        elif t == "ensemble":
            # 重唱：无边框表格，每个角色一列并排显示
            p = doc.add_paragraph()
            style_run(p.add_run("（重唱）"), "黑体", bold=True)
            children = e.get("children", [])
            if children:
                table = doc.add_table(rows=1, cols=len(children))
                for i, c in enumerate(children):
                    cp = table.rows[0].cells[i].paragraphs[0]
                    cname = c.get("characterName") or ""
                    if cname:
                        style_run(cp.add_run(cname), "黑体", bold=True, ascii_font="黑体")
                    style_run(cp.add_run("（唱）"), "黑体", bold=True, ascii_font="黑体").add_break()
                    lines = (c["content"] or "").split("\n")
                    for j, line in enumerate(lines):
                        run = style_run(cp.add_run(line), "仿宋")
                        if j < len(lines) - 1:
                            run.add_break()
        elif t == "dual":
            # 叠白：无边框表格，每个角色一列并排显示（同时对白）
            p = doc.add_paragraph()
            style_run(p.add_run("（叠白）"), "黑体", bold=True)
            children = e.get("children", [])
            if children:
                table = doc.add_table(rows=1, cols=len(children))
                for i, c in enumerate(children):
                    cp = table.rows[0].cells[i].paragraphs[0]
                    cname = c.get("characterName") or ""
                    if cname:
                        style_run(cp.add_run(cname), "黑体", bold=True, ascii_font="黑体").add_break()
                    lines = (c["content"] or "").split("\n")
                    for j, line in enumerate(lines):
                        run = style_run(cp.add_run(line), ea_body)
                        if j < len(lines) - 1:
                            run.add_break()
        elif t == "song":
            p = doc.add_paragraph()
            style_run(p.add_run(f"♪ {e['song_title'] or '（未命名歌曲）'}"), "黑体", bold=True)
            for c in e.get("children", []):
                add_script_element(c, in_song=True)
            # 歌曲结束标识
            end_p = doc.add_paragraph()
            style_run(end_p.add_run("（歌曲结束）"), "黑体", bold=True)
        else:
            p = doc.add_paragraph()
            add_text_with_breaks(p, e["content"], ea_body)

    doc = Document()
    style = doc.styles["Normal"]
    set_ea_style(style, ea_body)
    style.font.size = Pt(font_size)
    style.paragraph_format.line_spacing = line_spacing
    # 标题样式同样显式指定中文字体，避免主题字体回退到日文字体
    for hs in ["Heading 1", "Heading 2", "Heading 3", "Title"]:
        try:
            set_ea_style(doc.styles[hs], "黑体")
        except KeyError:
            pass

    default_name = "export"
    if chapter_id:
        data = _get_chapter_content(chapter_id)
        if not data: raise ValueError("Chapter not found")
        default_name = data["chapterTitle"] or default_name
        if data["volumeTitle"]:
            doc.add_heading(data["volumeTitle"], level=1)
        doc.add_heading(data["chapterTitle"], level=2)
        for para in (data["content"] or "").split("\n"):
            doc.add_paragraph(para)
    elif project_id:
        ptype, ptitle = _get_project_type(project_id)
        default_name = ptitle or default_name
        if ptype == "novel":
            volumes = _get_project_contents(project_id)
            for vol in volumes:
                doc.add_heading(vol["volumeTitle"], level=1)
                if vol["volumePreface"]:
                    # 斜体需设置在 run 上（python-docx 的 Paragraph 没有 italic 属性）
                    p = doc.add_paragraph()
                    style_run(p.add_run(vol["volumePreface"]), ea_body, italic=True)
                for ch in vol["chapters"]:
                    doc.add_heading(ch["title"], level=2)
                    for para in (ch["content"] or "").split("\n"):
                        doc.add_paragraph(para)
        else:
            for act in _get_script_contents(project_id):
                doc.add_heading(act["actTitle"], level=1)
                for sc in act["scenes"]:
                    doc.add_heading(sc["title"], level=2)
                    if sc["setting"]:
                        p = doc.add_paragraph()
                        style_run(p.add_run(f"（场景）{sc['setting']}"), "楷体", italic=True)
                    for e in _merge_action_dialogue(sc["elements"]):
                        add_script_element(e)

    out = _resolve_output_path(output_path, "docx", default_name)
    try:
        doc.save(out)
    except OSError as e:
        raise ValueError(f"无法写入文件 {out}：{e}")
    return {"filePath": out}
